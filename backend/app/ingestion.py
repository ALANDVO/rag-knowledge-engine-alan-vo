"""Document ingestion: PDF, DOCX, HTML, plain text with proper extraction."""
import structlog
from dataclasses import dataclass, field
from typing import Any
import io

log = structlog.get_logger()


@dataclass
class Document:
    doc_id: str
    title: str
    content: str
    source: str
    metadata: dict = field(default_factory=dict)


class IngestionPipeline:
    """Ingests documents from various formats with proper text extraction."""

    async def ingest(self, doc_id: str, title: str, content: str,
                     source: str, metadata: dict = None) -> Document:
        log.info("ingest_start", doc_id=doc_id, title=title, size=len(content))
        doc = Document(doc_id=doc_id, title=title, content=content,
                       source=source, metadata=metadata or {})
        log.info("ingest_done", doc_id=doc_id)
        return doc

    def _extract_pdf(self, data: bytes) -> str:
        """Extract text from PDF using pypdf."""
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            texts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    texts.append(text)
            return "\n\n".join(texts)
        except Exception as e:
            log.warning("pdf_extraction_failed", error=str(e))
            return data.decode("utf-8", errors="replace")

    def _extract_docx(self, data: bytes) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(io.BytesIO(data))
            texts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    texts.append(para.text)
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            texts.append(cell.text)
            return "\n\n".join(texts)
        except Exception as e:
            log.warning("docx_extraction_failed", error=str(e))
            return data.decode("utf-8", errors="replace")

    def _extract_html(self, data: bytes) -> str:
        """Extract text from HTML using BeautifulSoup."""
        try:
            from bs4 import BeautifulSoup
            html = data.decode("utf-8", errors="replace")
            soup = BeautifulSoup(html, "lxml")
            # Remove script and style elements
            for script in soup(["script", "style", "noscript"]):
                script.decompose()
            # Get text with reasonable spacing
            text = soup.get_text(separator="\n", strip=True)
            # Collapse multiple blank lines
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            return "\n\n".join(lines)
        except Exception as e:
            log.warning("html_extraction_failed", error=str(e))
            return data.decode("utf-8", errors="replace")

    async def ingest_file(self, doc_id: str, filename: str,
                          data: bytes) -> Document:
        if "." in filename:
            ext = filename.rsplit(".", 1)[-1].lower()
        else:
            ext = filename.lower()

        if ext == "pdf":
            content = self._extract_pdf(data)
            source = "file:pdf"
        elif ext == "docx":
            content = self._extract_docx(data)
            source = "file:docx"
        elif ext in ("html", "htm"):
            content = self._extract_html(data)
            source = "file:html"
        else:
            # Plain text, markdown, etc.
            content = data.decode("utf-8", errors="replace")
            source = f"file:{ext}"

        return await self.ingest(doc_id, filename, content, source)


ingestion = IngestionPipeline()